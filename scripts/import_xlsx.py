#!/usr/bin/env python3
"""Import gymnastics score sheets into src/data/competitionData.ts.

Downloads the Google Sheet (source of truth), caches it locally, then generates
the TypeScript dataset.
"""

from __future__ import annotations

import re
import urllib.request
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "resources" / "Gymnastics data.xlsx"
OUT = ROOT / "src" / "data" / "competitionData.ts"
GOOGLE_SHEET_ID = "1L63Vp5RpIcu5GQTE-hK2OhSfKNLfmh_9MEOYmy4xi_k"
GOOGLE_XLSX_URL = (
    f"https://docs.google.com/spreadsheets/d/{GOOGLE_SHEET_ID}/export?format=xlsx"
)

APPARATUS = ("vault", "bars", "beam", "floor")

INDIVIDUAL_SHEETS = [
    {
        "sheet": "Counties Manukau Comp 1 ",
        "id": "cm-1",
        "name": "Counties Manukau Comp 1",
    },
    {
        "sheet": "Counties Manukau Comp 2 (Waitak",
        "id": "cm-2",
        "name": "Counties Manukau Comp 2",
    },
    {
        "sheet": "Tristar comp 1 (Waitakere)",
        "id": "tristar-1",
        "name": "Tri Star Comp 1",
    },
    {
        "sheet": "Tristar comp 2",
        "id": "tristar-2",
        "name": "Tri Star Comp 2",
    },
    {
        "sheet": "CSG Comp 1",
        "id": "csg-1",
        "name": "CSG Comp 1",
    },
    {
        "sheet": "CSG Comp 2 (Waitakere)",
        "id": "csg-2",
        "name": "CSG Comp 2",
    },
    {
        "sheet": "NHG Comp",
        "id": "nhg-1",
        "name": "NHG Comp",
    },
    {
        "sheet": "Central Comp",
        "id": "central-1",
        "name": "Central Comp",
    },
    {
        "sheet": "Auckland Manukau Champs 1",
        "id": "amc-1",
        "name": "Auckland Manukau Champs 1",
    },
    {
        "sheet": "Auckland Manukau Champs 2 (Wait",
        "id": "amc-2",
        "name": "Auckland Manukau Champs 2",
    },
]

TEAM_SHEETS = [
    {"sheet": "Teams Counties Manukau Comp 1", "competitionId": "cm-1"},
    {"sheet": "Teams Counties Manukau Comp 2 (", "competitionId": "cm-2"},
    {"sheet": "Teams Tristar Comp 1 (Waitakere", "competitionId": "tristar-1"},
    {"sheet": "Teams Tristar Comp 2", "competitionId": "tristar-2"},
    {"sheet": "Teams NHG Comp", "competitionId": "nhg-1"},
    {"sheet": "Teams Central Comp", "competitionId": "central-1"},
    {"sheet": "Teams Auckland Manukau Champs C", "competitionId": "amc-1"},
    {"sheet": "Sheet13", "competitionId": "amc-2"},
]

CLUB_ALIASES = {
    "nhg": "North Harbour Gymnastics",
    "nhg gymnastics": "North Harbour Gymnastics",
    "turn & gymnastic circle": "Turn and Gymnastic Circle",
    "turn & gymnastics circle": "Turn and Gymnastic Circle",
    "turn and gymnastic circle": "Turn and Gymnastic Circle",
    "howick gymnastic club": "Howick Gymnastics",
    "howick gymnastics club": "Howick Gymnastics",
    "howick gymnastics": "Howick Gymnastics",
    "eastern suburbs gymnastic club": "Eastern Suburbs Gymnastics",
    "eastern suburbs gymnastics club": "Eastern Suburbs Gymnastics",
    "eastern suburbs gym club": "Eastern Suburbs Gymnastics",
    "eastern suburbs gymnastics": "Eastern Suburbs Gymnastics",
    "mid-island gym sports": "Mid-Island GymSports",
    "mid-island gymsports": "Mid-Island GymSports",
    "impact gymnastics academy": "Impact Gymsport Academy",
    "argos gymnastics": "ARGOS Gymnastics Club",
    "argos gymnastics club": "ARGOS Gymnastics Club",
    "olympia gymnastic sports": "Olympia Gymsports",
    "olympia gymsports": "Olympia Gymsports",
    "olympia gymnastics": "Olympia Gymsports",
    "te wero": "Te Wero Gymnastics",
    "te wero gymnastics": "Te Wero Gymnastics",
    "tristar gymnastics": "Tri Star Gymnastics",
    "tri star gymnastics": "Tri Star Gymnastics",
    "rimutaka gymnastics": "Rimutaka Gymsports",
    "nelson gymnastics": "Gymnastics Nelson",
    "blenheim gymnastics": "Blenheim Gymnastics Club",
    "impact gymnastics": "Impact Gymsport Academy",
    "omni gymnastics": "Omni Gymnastics Centre",
    "manawatu gymnastics incorporated": "Manawatu Gymnastics",
    "manawatu gymnastics": "Manawatu Gymnastics",
    "bay of islands gymnastics": "Bay of Islands Gymnastics",
    "bay of islands gymnastics club": "Bay of Islands Gymnastics",
    "christchurch school of gymnastics": "Christchurch School of Gymnastics",
    "whangarei academy of gymnastics": "Whangarei Academy of Gymnastics",
    "mt tauhara gymnastics club": "Mt Tauhara Gymnastics Club",
    "hamilton city gymnastics": "Hamilton City Gymnastics",
}

CLUB_META = {
    "Turn and Gymnastic Circle": ("TAG", "#5b50e6", "Waikato"),
    "Whangarei Academy of Gymnastics": ("WAO", "#ed6a5a", "Northland"),
    "Howick Gymnastics": ("HG", "#1a9c78", "Howick"),
    "North Harbour Gymnastics": ("NHG", "#d89b2b", "North Shore"),
    "Mid-Island GymSports": ("MIG", "#2f6fed", "Bay of Plenty"),
    "Counties Manukau Gymnastics": ("CMG", "#c44d8a", "Manukau"),
    "Gymnastics Waitara": ("GW", "#0f766e", "Taranaki"),
    "Waitakere Gymnastics": ("WG", "#b45309", "Auckland"),
    "Tri Star Gymnastics": ("TSG", "#7c3aed", "Auckland"),
    "Eastern Suburbs Gymnastics": ("ESG", "#be123c", "Auckland"),
    "Impact Gymsport Academy": ("IGA", "#0369a1", "Bay of Plenty"),
    "ARGOS Gymnastics Club": ("AGC", "#4d7c0f", "Bay of Plenty"),
    "Bay of Islands Gymnastics": ("BOI", "#9333ea", "Northland"),
    "Affinity Gymnastics Academy": ("AGA", "#c2410c", "Canterbury"),
    "Christchurch School of Gymnastics": ("CSG", "#0e7490", "Christchurch"),
    "Hutt Valley Gymnastics": ("HVG", "#a16207", "Wellington"),
    "Gymnastics Nelson": ("GN", "#db2777", "Nelson"),
    "Dunedin Gymnastics Academy": ("DGA", "#15803d", "Otago"),
    "Te Wero Gymnastics": ("TWG", "#4338ca", "Waikato"),
    "Olympia Gymsports": ("OG", "#ea580c", "Canterbury"),
    "Blenheim Gymnastics Club": ("BGC", "#0891b2", "Marlborough"),
    "Invercargill Gymnastics Club": ("IGC", "#65a30d", "Southland"),
    "Rimutaka Gymsports": ("RG", "#e11d48", "Upper Hutt"),
    "Mt Tauhara Gymnastics Club": ("MTG", "#2563eb", "Waikato"),
    "Hamilton City Gymnastics": ("HCG", "#7c2d12", "Hamilton"),
    "Twisters Gymnastics": ("TW", "#c026d3", "Wellington"),
    "Whanganui Boys and Girls Gym Club": ("WBG", "#b45309", "Whanganui"),
    "Levin Gymnastics": ("LG", "#0d9488", "Levin"),
    "Omni Gymnastics Centre": ("OGC", "#be123c", "Hawke's Bay"),
    "Kapiti Gymnastics": ("KG", "#1d4ed8", "Kapiti"),
    "Manawatu Gymnastics": ("MG", "#4d7c0f", "Manawatū"),
}

CITY_FIXES = {
    "lower hut": "Lower Hutt",
}

EXTRA_COLORS = [
    "#0f766e",
    "#7c3aed",
    "#b45309",
    "#0369a1",
    "#be123c",
    "#4d7c0f",
    "#9333ea",
    "#c2410c",
]


def slugify(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    return value.strip("-")


def excel_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, (int, float)) and value > 30000:
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
    return None


def format_location(address: str | None) -> str:
    """Turn a full address into 'Neighborhood, City' (drop street + postcode)."""
    if not address:
        return ""
    parts = [part.strip() for part in str(address).split(",") if part.strip()]
    if len(parts) >= 2:
        city = re.sub(r"\s+\d+\s*$", "", parts[-1]).strip()
        city = CITY_FIXES.get(city.lower(), city)
        neighborhood = parts[-2]
        if city:
            return f"{neighborhood}, {city}"
        return neighborhood
    return re.sub(r"\s+\d+\s*$", "", parts[0]).strip() if parts else ""


def find_header_row(rows: list) -> int:
    for index, row in enumerate(rows[:10]):
        if not row:
            continue
        cells = [str(cell).strip().lower() if cell is not None else "" for cell in row[:5]]
        if "rank" in cells and "name" in cells:
            return index
    return 3


def sheet_date(rows: list) -> str | None:
    for row in rows[:5]:
        if not row:
            continue
        for cell in row:
            date = excel_date(cell)
            if date:
                return date
    return None


def sheet_location(rows: list, header_index: int) -> str:
    for row in rows[:header_index]:
        if not row or row[0] is None:
            continue
        if excel_date(row[0]):
            continue
        text = str(row[0]).strip()
        if "," in text:
            return format_location(text)
    return ""


def parse_score(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip()
        if not text or text.upper() in {"DNS", "DNF", "-"}:
            return None
        text = text.rstrip("*").strip()
        try:
            return float(text)
        except ValueError:
            return None
    return None


def normalize_club(raw: str | None) -> str:
    if not raw:
        return "Unknown Club"
    name = " ".join(str(raw).split())
    return CLUB_ALIASES.get(name.lower(), name)


def short_name_for(name: str) -> str:
    if name in CLUB_META:
        return CLUB_META[name][0]
    parts = re.findall(r"[A-Za-z0-9]+", name)
    if not parts:
        return name[:3].upper()
    if len(parts) == 1:
        return parts[0][:3].upper()
    return "".join(part[0] for part in parts[:3]).upper()


def color_for(name: str, index: int) -> str:
    if name in CLUB_META:
        return CLUB_META[name][1]
    return EXTRA_COLORS[index % len(EXTRA_COLORS)]


def city_for(name: str) -> str:
    if name in CLUB_META:
        return CLUB_META[name][2]
    return ""


def ts_string(value: str) -> str:
    return "'" + value.replace("\\", "\\\\").replace("'", "\\'") + "'"


def download_google_sheet() -> None:
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    print(f"Downloading Google Sheet {GOOGLE_SHEET_ID}…")
    request = urllib.request.Request(
        GOOGLE_XLSX_URL,
        headers={"User-Agent": "Mozilla/5.0 (Gymnastics importer)"},
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        payload = response.read()
    if len(payload) < 1000 or not payload.startswith(b"PK"):
        raise RuntimeError(
            "Google Sheet export did not return an xlsx file. "
            "Check that the sheet is shared as 'Anyone with the link'."
        )
    XLSX.write_bytes(payload)
    print(f"Saved {XLSX.relative_to(ROOT)} ({len(payload)} bytes)")


def main() -> None:
    download_google_sheet()
    wb = load_workbook(XLSX, data_only=True)
    sheet_by_prefix = {name: name for name in wb.sheetnames}

    def resolve_sheet(prefix: str) -> str:
        if prefix in sheet_by_prefix:
            return prefix
        matches = [name for name in wb.sheetnames if name.startswith(prefix.rstrip())]
        if not matches:
            raise KeyError(f"Missing sheet starting with {prefix!r}")
        return matches[0]

    clubs: dict[str, dict] = {}
    gymnasts_by_name: dict[str, dict] = {}
    scores: list[dict] = []
    competitions: list[dict] = []
    competition_teams: list[dict] = []

    class State:
        next_id = 1

    def ensure_club(raw_name: str | None) -> str:
        name = normalize_club(raw_name)
        club_id = slugify(name)
        if club_id not in clubs:
            clubs[club_id] = {
                "id": club_id,
                "name": name,
                "shortName": short_name_for(name),
                "color": color_for(name, len(clubs)),
                "city": city_for(name),
            }
        return club_id

    def ensure_gymnast(name: str, club_id: str) -> str:
        key = name.strip().lower()
        existing = gymnasts_by_name.get(key)
        if existing:
            if existing["teamId"] == "unknown-club" and club_id != "unknown-club":
                existing["teamId"] = club_id
            return existing["id"]
        gymnast_id = f"g{State.next_id}"
        State.next_id += 1
        gymnasts_by_name[key] = {
            "id": gymnast_id,
            "name": " ".join(name.split()),
            "teamId": club_id,
        }
        return gymnast_id

    # Individual sheets
    for meta in INDIVIDUAL_SHEETS:
        sheet_name = resolve_sheet(meta["sheet"])
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        header_index = find_header_row(rows)
        date = sheet_date(rows)
        location = sheet_location(rows, header_index)
        competitions.append(
            {
                "id": meta["id"],
                "name": meta["name"],
                "date": date or "2026-01-01",
                "location": location,
                "category": "Women",
                "level": "Open",
            }
        )

        for row in rows[header_index + 1 :]:
            if not row or row[1] is None:
                continue
            name = str(row[1]).strip()
            if not name:
                continue
            club_id = ensure_club(row[2])
            gymnast_id = ensure_gymnast(name, club_id)
            apparatus_scores = {
                "vault": parse_score(row[4]),
                "bars": parse_score(row[6]),
                "beam": parse_score(row[8]),
                "floor": parse_score(row[10]),
            }
            # Skip full DNS rows
            if all(value is None for value in apparatus_scores.values()):
                continue
            for apparatus, score in apparatus_scores.items():
                if score is None:
                    continue
                scores.append(
                    {
                        "competitionId": meta["id"],
                        "gymnastId": gymnast_id,
                        "apparatus": apparatus,
                        "score": score,
                    }
                )

    # Team sheets
    for meta in TEAM_SHEETS:
        sheet_name = resolve_sheet(meta["sheet"])
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        current_team = None

        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            kind = str(row[0]).strip()
            if kind == "Team":
                team_name = str(row[2]).strip()
                club_raw = row[4]
                club_id = None
                # Mix / multi-club teams keep no single club id
                if club_raw and "/" not in str(club_raw):
                    club_id = ensure_club(club_raw)
                color = (
                    clubs[club_id]["color"]
                    if club_id
                    else color_for(team_name, len(competition_teams))
                )
                short = short_name_for(team_name)
                team_id = f"{meta['competitionId']}-{slugify(team_name)}"
                current_team = {
                    "id": team_id,
                    "competitionId": meta["competitionId"],
                    "name": team_name,
                    "shortName": short,
                    "clubId": club_id,
                    "color": color,
                    "rank": int(row[1]) if isinstance(row[1], (int, float)) else 0,
                    "apparatus": {
                        "vault": parse_score(row[5]) or 0.0,
                        "bars": parse_score(row[7]) or 0.0,
                        "beam": parse_score(row[9]) or 0.0,
                        "floor": parse_score(row[11]) or 0.0,
                    },
                    "total": parse_score(row[13]) or 0.0,
                    "members": [],
                }
                competition_teams.append(current_team)
            elif kind == "Gymnast" and current_team is not None:
                name = str(row[3]).strip()
                club_id = ensure_club(row[4])
                gymnast_id = ensure_gymnast(name, club_id)
                member_scores = {
                    "vault": parse_score(row[5]),
                    "bars": parse_score(row[7]),
                    "beam": parse_score(row[9]),
                    "floor": parse_score(row[11]),
                }
                total = parse_score(row[13])
                current_team["members"].append(
                    {
                        "gymnastId": gymnast_id,
                        "scores": {
                            key: value
                            for key, value in member_scores.items()
                            if value is not None
                        },
                        "total": total,
                    }
                )

    gymnast_list = list(gymnasts_by_name.values())
    gymnast_list.sort(key=lambda item: int(item["id"][1:]))
    team_list = sorted(clubs.values(), key=lambda item: item["name"].lower())

    # Emit TypeScript
    lines: list[str] = []
    lines.append("import type { CompetitionDataset } from './types'")
    lines.append("")
    lines.append(
        "/** Imported from the Gymnastics Google Sheet — individual + official team results */"
    )
    lines.append("export const dataset: CompetitionDataset = {")

    lines.append("  competitions: [")
    for item in competitions:
        lines.append(
            "    {"
            f" id: {ts_string(item['id'])},"
            f" name: {ts_string(item['name'])},"
            f" date: {ts_string(item['date'])},"
            f" location: {ts_string(item['location'])},"
            f" category: {ts_string(item['category'])},"
            f" level: {ts_string(item['level'])},"
            " },"
        )
    lines.append("  ],")

    lines.append("  teams: [")
    for item in team_list:
        lines.append(
            "    {"
            f" id: {ts_string(item['id'])},"
            f" name: {ts_string(item['name'])},"
            f" shortName: {ts_string(item['shortName'])},"
            f" color: {ts_string(item['color'])},"
            f" city: {ts_string(item['city'])},"
            " },"
        )
    lines.append("  ],")

    lines.append("  gymnasts: [")
    for item in gymnast_list:
        lines.append(
            "    {"
            f" id: {ts_string(item['id'])},"
            f" name: {ts_string(item['name'])},"
            f" teamId: {ts_string(item['teamId'])},"
            " },"
        )
    lines.append("  ],")

    lines.append("  scores: [")
    for item in scores:
        score = item["score"]
        score_text = (
            str(int(score))
            if isinstance(score, float) and score.is_integer()
            else repr(score)
        )
        lines.append(
            "    {"
            f" competitionId: {ts_string(item['competitionId'])},"
            f" gymnastId: {ts_string(item['gymnastId'])},"
            f" apparatus: {ts_string(item['apparatus'])},"
            f" score: {score_text},"
            " },"
        )
    lines.append("  ],")

    lines.append("  competitionTeams: [")
    for team in competition_teams:
        lines.append("    {")
        lines.append(f"      id: {ts_string(team['id'])},")
        lines.append(f"      competitionId: {ts_string(team['competitionId'])},")
        lines.append(f"      name: {ts_string(team['name'])},")
        lines.append(f"      shortName: {ts_string(team['shortName'])},")
        if team["clubId"]:
            lines.append(f"      clubId: {ts_string(team['clubId'])},")
        lines.append(f"      color: {ts_string(team['color'])},")
        lines.append(f"      rank: {team['rank']},")
        lines.append(f"      total: {repr(team['total'])},")
        app = team["apparatus"]
        lines.append(
            "      apparatus: {"
            f" vault: {repr(app['vault'])},"
            f" bars: {repr(app['bars'])},"
            f" beam: {repr(app['beam'])},"
            f" floor: {repr(app['floor'])},"
            " },"
        )
        lines.append("      members: [")
        for member in team["members"]:
            score_parts = ", ".join(
                f"{key}: {repr(value)}" for key, value in member["scores"].items()
            )
            total_part = (
                "null" if member["total"] is None else repr(member["total"])
            )
            lines.append(
                "        {"
                f" gymnastId: {ts_string(member['gymnastId'])},"
                f" scores: {{ {score_parts} }},"
                f" total: {total_part},"
                " },"
            )
        lines.append("      ],")
        lines.append("    },")
    lines.append("  ],")
    lines.append("}")
    lines.append("")

    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(
        f"Wrote {OUT.relative_to(ROOT)}: "
        f"{len(competitions)} comps, {len(team_list)} clubs, "
        f"{len(gymnast_list)} gymnasts, {len(scores)} scores, "
        f"{len(competition_teams)} competition teams"
    )


if __name__ == "__main__":
    main()
